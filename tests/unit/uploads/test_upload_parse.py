"""Tests for avtalsagent.uploads.parse, extract and sections: a file into sections, or refused.

A PDF, a Word file and text files are cut at their numbered headings, with
the page each section starts on; headings without numbers ("§ 3", "Bilaga
2", Word's and Markdown's) are used when there is no numbered outline; a
long section is cut into pieces that keep their pages. Each limit and each
file that cannot be read gives its status and Swedish text. The PDFs and
Word files are built in the tests (`upload_files.py`).
"""

import pytest
from docx.document import Document as WordDocument

from avtalsagent.domain.uploads import UploadKind
from avtalsagent.uploads.errors import (
    EMPTY,
    NO_TEXT,
    NOT_UTF8,
    SCANNED,
    UNREADABLE,
    UploadRejected,
)
from avtalsagent.uploads.parse import ParsedFile, UploadLimits, parse_upload
from avtalsagent.uploads.sections import MAX_SECTION_CHARS
from tests.unit.uploads.upload_files import AGREEMENT_PAGES, docx_bytes, pdf_bytes, zip_bytes

LIMITS = UploadLimits(max_bytes=10 * 1024 * 1024, max_pages=300, max_characters=1_500_000)
FILLER = "Leverantören ska utföra tjänsten enligt avtalet och dess bilagor med omsorg."


def parse(data: bytes, filename: str, limits: UploadLimits = LIMITS) -> ParsedFile:
    return parse_upload(data, filename, limits)


def rejected(data: bytes, filename: str, limits: UploadLimits = LIMITS) -> UploadRejected:
    with pytest.raises(UploadRejected) as caught:
        parse(data, filename, limits)
    return caught.value


def headings(parsed: ParsedFile) -> list[tuple[str | None, str, int | None]]:
    return [(s.number, s.title, s.page_start) for s in parsed.sections]


def test_a_pdf_is_cut_at_its_numbered_headings_with_the_page_each_starts_on() -> None:
    parsed = parse(pdf_bytes(AGREEMENT_PAGES), "avtal.pdf")

    assert parsed.kind is UploadKind.PDF
    assert parsed.pages == 4
    assert headings(parsed) == [
        (None, "Text före första rubriken", 1),
        ("1", "Parter", 1),
        ("2", "Avtalstid", 1),
        ("2.1", "Förlängning", 2),
        ("3", "Uppsägning", 4),
    ]
    assert [s.position for s in parsed.sections] == [0, 1, 2, 3, 4]
    assert [s.level for s in parsed.sections] == [0, 1, 1, 2, 1]
    # A PDF's lines are joined by single line breaks, also across a page.
    assert parsed.sections[2].text == "2 Avtalstid\nAvtalet gäller i två år."
    assert parsed.characters == sum(len(s.text) for s in parsed.sections)


def test_pages_without_text_are_named_in_a_warning() -> None:
    parsed = parse(pdf_bytes(AGREEMENT_PAGES), "avtal.pdf")

    assert parsed.warnings == (
        "Sidan 3 har ingen text som går att läsa (den kan vara inskannad). "
        "Agenten ser inte vad som står på den.",
    )
    pages = [*AGREEMENT_PAGES, [], []]
    (warning,) = parse(pdf_bytes(pages), "avtal.pdf").warnings
    assert warning.startswith("Sidorna 3, 5 och 6 har ingen text")


def test_a_pdf_without_text_is_refused_as_scanned() -> None:
    error = rejected(pdf_bytes([[], []]), "inskannad.pdf")

    assert (error.status_code, error.detail) == (422, SCANNED)


def test_a_page_with_only_a_stamped_page_number_has_no_text() -> None:
    error = rejected(pdf_bytes([["1"], ["2"]]), "inskannad.pdf")

    assert error.status_code == 422


def test_a_pdf_with_too_many_pages_is_413_before_its_text_is_read() -> None:
    limits = UploadLimits(max_bytes=LIMITS.max_bytes, max_pages=3, max_characters=1_000)

    error = rejected(pdf_bytes(AGREEMENT_PAGES), "avtal.pdf", limits)

    assert error.status_code == 413
    assert error.detail == "PDF:en har 4 sidor. Den får ha högst 3."


def test_a_broken_pdf_is_422() -> None:
    error = rejected(b"%PDF-1.7\nnot really a pdf", "avtal.pdf")

    assert (error.status_code, error.detail) == (422, UNREADABLE)


def test_too_much_text_is_413() -> None:
    limits = UploadLimits(max_bytes=LIMITS.max_bytes, max_pages=300, max_characters=100)

    error = rejected(pdf_bytes(AGREEMENT_PAGES), "avtal.pdf", limits)

    assert error.status_code == 413
    assert "högst 100" in error.detail


def test_too_many_bytes_and_an_empty_file_are_refused() -> None:
    limits = UploadLimits(max_bytes=1024 * 1024, max_pages=300, max_characters=1_000)

    error = rejected(b"a" * (1024 * 1024 + 1), "noter.txt", limits)
    assert (error.status_code, error.detail) == (413, "Filen är för stor. Den får vara högst 1 MB.")
    error = rejected(b"", "noter.txt")
    assert (error.status_code, error.detail) == (422, EMPTY)


def test_a_word_file_is_cut_at_its_headings_and_keeps_its_tables() -> None:
    def build(document: WordDocument) -> None:
        document.add_heading("Allmänna villkor", 0)
        document.add_heading("1 Inledning", 1)
        document.add_paragraph("Detta avtal gäller.")
        document.add_heading("2 Ansvar", 1)
        document.add_paragraph("Leverantören ansvarar.", style="List Bullet")
        document.add_heading("2.1 Begränsning", 2)
        table = document.add_table(rows=2, cols=2)
        for row, cells in enumerate([("Roll", "Pris"), ("Konsult", "1 000 kr")]):
            for column, text in enumerate(cells):
                table.cell(row, column).text = text

    parsed = parse(docx_bytes(build), "villkor.docx")

    assert parsed.kind is UploadKind.DOCX
    assert parsed.pages is None
    assert headings(parsed) == [
        (None, "Allmänna villkor", None),
        ("1", "Inledning", None),
        ("2", "Ansvar", None),
        ("2.1", "Begränsning", None),
    ]
    assert parsed.sections[3].text == "2.1 Begränsning\n\nRoll | Pris\nKonsult | 1 000 kr"


def test_a_word_file_without_numbers_is_cut_at_its_heading_styles() -> None:
    def build(document: WordDocument) -> None:
        for title in ("Bakgrund", "Pris", "Ansvar"):
            document.add_heading(title, 1)
            document.add_paragraph(f"Om {title.lower()}.")

    parsed = parse(docx_bytes(build), "villkor.docx")

    assert headings(parsed) == [
        (None, "Bakgrund", None),
        (None, "Pris", None),
        (None, "Ansvar", None),
    ]


def test_a_broken_word_file_is_422_and_logged(caplog: pytest.LogCaptureFixture) -> None:
    broken = zip_bytes({"word/document.xml": b"<not xml", "[Content_Types].xml": b"<x"})

    error = rejected(broken, "avtal.docx")

    assert (error.status_code, error.detail) == (422, UNREADABLE)
    assert "could not read the uploaded docx file" in caplog.text


def test_an_empty_word_file_has_no_text() -> None:
    error = rejected(docx_bytes(lambda document: None), "tom.docx")

    assert (error.status_code, error.detail) == (422, NO_TEXT)


def test_markdown_headings_and_numbered_lines_cut_a_text_file() -> None:
    text = "# Mitt avtal\n\nInledning.\n\n## 1 Parter\n\nA och B.\n\n## 2 Pris\nPriset är 5 kr.\n"

    parsed = parse(text.encode(), "avtal.md")

    assert parsed.kind is UploadKind.TEXT
    assert headings(parsed)[1:] == [("1", "Parter", None), ("2", "Pris", None)]
    assert parsed.sections[2].text == "2 Pris\n\nPriset är 5 kr."


def test_paragraphs_and_bilagor_cut_a_text_file_without_numbers() -> None:
    text = "§ 1 Inledning\nText ett.\n\n§ 2 Pris\nText två,\npå två rader.\n\nBilaga 1 Lista\nRad."

    parsed = parse(text.encode(), "avtal.txt")

    assert headings(parsed) == [
        (None, "§ 1 Inledning", None),
        (None, "§ 2 Pris", None),
        (None, "Bilaga 1 Lista", None),
    ]
    assert parsed.sections[1].text == "§ 2 Pris\n\nText två,\npå två rader."


def test_a_wrapped_line_about_a_bilaga_is_not_a_heading() -> None:
    text = "Inledning\n\nPriserna står i\nBilaga 2 och gäller ett år."

    parsed = parse(text.encode(), "avtal.txt")

    assert len(parsed.sections) == 1


def test_text_must_be_utf8_and_may_have_a_byte_order_mark() -> None:
    error = rejected("Sjöräddning".encode("latin-1"), "noter.txt")
    assert (error.status_code, error.detail) == (422, NOT_UTF8)

    parsed = parse("﻿Sjöräddning".encode(), "noter.txt")
    assert parsed.sections[0].text == "Sjöräddning"


def test_whitespace_alone_is_no_text() -> None:
    error = rejected(b" \n\n \t\n", "noter.txt")

    assert (error.status_code, error.detail) == (422, NO_TEXT)


def test_a_long_section_is_cut_into_pieces_that_keep_their_pages() -> None:
    # Twelve pages of 50 lines without headings: one section of about 55,000 characters.
    pages = [[f"{FILLER} Sida {page}, rad {row}." for row in range(50)] for page in range(1, 13)]

    parsed = parse(pdf_bytes(pages), "lång.pdf")

    pieces = parsed.sections
    assert len(pieces) == 5
    assert all(len(piece.text) <= MAX_SECTION_CHARS for piece in pieces)
    assert [piece.title for piece in pieces] == [
        f"Text före första rubriken (del {n} av 5)" for n in range(1, 6)
    ]
    for piece in pieces:
        # Each piece starts with a whole line, on the page that line is on.
        first = piece.text.split("\n", 1)[0]
        assert first.startswith(FILLER)
        assert f"Sida {piece.page_start}," in first
    assert pieces[0].page_start == 1
    assert [piece.page_start for piece in pieces] == sorted(p.page_start or 0 for p in pieces)
    joined = "\n".join(piece.text for piece in pieces)
    assert joined.count(FILLER) == 12 * 50  # no line lost or doubled
