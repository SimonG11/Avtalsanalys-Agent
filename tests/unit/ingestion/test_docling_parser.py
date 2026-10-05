"""Tests for avtalsagent.ingestion.parsers.docling_parser.

Word files need no model. PDFs need Docling's layout and table models, which are
downloaded from Hugging Face on first use; where they cannot be downloaded the PDF
tests are skipped with the reason, unless AVTALSAGENT_REQUIRE_MODELS=1 (set in CI),
which makes them fail instead.
"""

import os
from pathlib import Path

import pytest
from docling.models.stages.page_assemble.page_assemble_model import PageAssembleOptions
from docling_core.types.doc import ContentLayer, DocItemLabel, DoclingDocument, TableCell, TableData
from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from avtalsagent.domain.parsed import Block, BlockKind
from avtalsagent.ingestion.parsers.docling_parser import (
    DoclingParser,
    _blocks,
    _KeepHyphensAssembleModel,
)


@pytest.fixture(scope="module")
def parser() -> DoclingParser:
    return DoclingParser()


def make_docx(target: Path) -> Path:
    document = Document()
    document.add_heading("1 Personuppgiftsbiträdesavtalets syfte", level=1)
    document.add_heading("1.1 Detta personuppgiftsbiträdesavtal reglerar behandlingen.", level=2)
    document.add_paragraph("Personuppgiftsansvarig modifierar och anpassar utkastet.")
    document.add_paragraph("Behandlingen avser kunddata", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Personuppgiftsansvarig:"
    table.cell(0, 1).text = "Myndigheten"
    table.cell(1, 0).text = "Organisationsnummer:"
    table.cell(1, 1).text = "202100-0829"
    document.add_heading("2 Parter", level=1)
    document.save(str(target))
    return target


def test_word_headings_keep_numbers_and_levels(parser: DoclingParser, tmp_path: Path) -> None:
    blocks = parser.parse(make_docx(tmp_path / "pub.docx"))
    headings = [(b.text, b.level) for b in blocks if b.kind is BlockKind.HEADING]
    assert headings == [
        ("1 Personuppgiftsbiträdesavtalets syfte", 1),
        ("1.1 Detta personuppgiftsbiträdesavtal reglerar behandlingen.", 2),
        ("2 Parter", 1),
    ]
    assert all(b.page is None for b in blocks)


def test_a_word_table_is_one_block_with_one_line_per_row(
    parser: DoclingParser, tmp_path: Path
) -> None:
    blocks = parser.parse(make_docx(tmp_path / "pub.docx"))
    tables = [b.text for b in blocks if b.kind is BlockKind.TABLE]
    assert tables == ["Personuppgiftsansvarig: | Myndigheten\nOrganisationsnummer: | 202100-0829"]
    # The cell text is not repeated as paragraphs after the table.
    assert not any(b.text == "Myndigheten" for b in blocks if b.kind is not BlockKind.TABLE)


def test_word_text_and_list_items(parser: DoclingParser, tmp_path: Path) -> None:
    blocks = parser.parse(make_docx(tmp_path / "pub.docx"))
    kinds = {b.text: b.kind for b in blocks}
    assert kinds["Personuppgiftsansvarig modifierar och anpassar utkastet."] is BlockKind.TEXT
    assert any(
        b.kind is BlockKind.LIST_ITEM and b.text.endswith("Behandlingen avser kunddata")
        for b in blocks
    )


def make_pdf(target: Path) -> Path:
    canvas = Canvas(str(target), pagesize=A4)
    _, height = A4
    for page in (1, 2):
        canvas.setFont("Helvetica-Bold", 14)
        canvas.drawString(72, height - 72, f"6.{page} Försäkring och ansvar del {page}")
        canvas.setFont("Helvetica", 11)
        for line in range(12):
            canvas.drawString(
                72,
                height - 110 - line * 16,
                f"Ramavtalsleverantören ska ha en ansvarsförsäkring, stycke {page}.{line}.",
            )
        canvas.setFont("Helvetica", 8)
        canvas.drawString(72, 40, f"Sida {page} (2)")
        canvas.showPage()
    canvas.save()
    return target


def make_price_pdf(target: Path) -> Path:
    styles = getSampleStyleSheet()
    rows = [
        ["Kompetensnivå", "Takpris per timme"],
        ["Nivå 1", "650 kr"],
        ["Nivå 2", "850 kr"],
        ["Nivå 3", "1 150 kr"],
    ]
    table = Table(rows, colWidths=[200, 200])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    SimpleDocTemplate(str(target), pagesize=A4).build(
        [
            Paragraph("3.2 Takpriser", styles["Heading2"]),
            Paragraph(
                "Ramavtalsleverantören får inte debitera mer än takpriset.", styles["Normal"]
            ),
            Spacer(1, 12),
            table,
        ]
    )
    return target


def parse_or_skip(parser: DoclingParser, path: Path) -> list[Block]:
    try:
        return parser.parse(path)
    except Exception as error:  # the models could not be downloaded
        if os.environ.get("AVTALSAGENT_REQUIRE_MODELS") == "1":
            raise
        pytest.skip(f"Docling's models are not available: {error}")


@pytest.fixture(scope="module")
def pdf_blocks(parser: DoclingParser, tmp_path_factory: pytest.TempPathFactory) -> list[Block]:
    return parse_or_skip(parser, make_pdf(tmp_path_factory.mktemp("pdf") / "terms.pdf"))


def test_pdf_text_comes_from_the_text_layer_with_pages(pdf_blocks: list[Block]) -> None:
    texts = [b.text for b in pdf_blocks]
    assert any("6.1 Försäkring och ansvar del 1" in text for text in texts)
    assert any("stycke 2.11." in text for text in texts)
    pages = {b.page for b in pdf_blocks if "stycke 2." in b.text}
    assert pages == {2}


def test_pdf_page_numbers_are_not_body_text(pdf_blocks: list[Block]) -> None:
    page_numbers = [b for b in pdf_blocks if b.text.startswith("Sida ")]
    assert all(b.kind is BlockKind.PAGE_FOOTER for b in page_numbers)


def test_a_pdf_table_gets_its_cells_from_the_table_model(
    parser: DoclingParser, tmp_path: Path
) -> None:
    blocks = parse_or_skip(parser, make_price_pdf(tmp_path / "prices.pdf"))
    tables = [b.text for b in blocks if b.kind is BlockKind.TABLE]
    assert tables == [
        "Kompetensnivå | Takpris per timme\nNivå 1 | 650 kr\nNivå 2 | 850 kr\nNivå 3 | 1 150 kr"
    ]
    assert any(b.kind is BlockKind.HEADING and b.text == "3.2 Takpriser" for b in blocks)


# How Docling's items become blocks, tested on documents built in code (no model).


def test_a_hyphen_at_a_line_break_is_kept() -> None:
    assemble = _KeepHyphensAssembleModel(options=PageAssembleOptions())
    lines = ["Ramavtalet är giltigt från och med 2025-", "08-19 för kompetensnivå 1-", "4."]
    assert assemble.sanitize_text(lines) == (
        "Ramavtalet är giltigt från och med 2025-08-19 för kompetensnivå 1-4."
    )
    assert assemble.sanitize_text(["Leverantören ska", "- föra en förteckning"]) == (
        "Leverantören ska - föra en förteckning"
    )


def test_text_inside_a_picture_is_kept() -> None:
    # TendSign draws its question boxes, which the layout model calls pictures.
    document = DoclingDocument(name="upphandlingsdokument")
    box = document.add_picture()
    document.add_text(DocItemLabel.TEXT, "Accepterar anbudsgivaren villkoren?", parent=box)
    document.add_text(DocItemLabel.TEXT, "Ja/Nej. Ja krävs", parent=box)
    assert [b.text for b in _blocks(document, None)] == [
        "Accepterar anbudsgivaren villkoren?",
        "Ja/Nej. Ja krävs",
    ]


def test_table_cells_with_the_same_text_and_the_footnote_are_kept() -> None:
    rows = [
        ["Onlinetjänst", "SSAE 18 SOC 1 Typ II", "SSAE 18 SOC 2 Typ II"],
        ["Office 365-tjänster", "Ja", "Ja"],
        ["Microsoft Azure Core Services", "", "Varierar*"],
    ]
    cells = [
        TableCell(
            text=text,
            start_row_offset_idx=r,
            end_row_offset_idx=r + 1,
            start_col_offset_idx=c,
            end_col_offset_idx=c + 1,
        )
        for r, row in enumerate(rows)
        for c, text in enumerate(row)
    ]
    # A heading row that spans all three columns appears in the grid once per column.
    title = TableCell(
        text="Certifieringar",
        col_span=3,
        start_row_offset_idx=0,
        end_row_offset_idx=1,
        start_col_offset_idx=0,
        end_col_offset_idx=3,
    )
    cells = [title] + [
        cell.model_copy(
            update={
                "start_row_offset_idx": cell.start_row_offset_idx + 1,
                "end_row_offset_idx": cell.end_row_offset_idx + 1,
            }
        )
        for cell in cells
    ]
    document = DoclingDocument(name="produktvillkor")
    table = document.add_table(data=TableData(num_rows=4, num_cols=3, table_cells=cells))
    note = "*Aktuell omfattning beskrivs i granskningsrapporten."
    document.add_text(DocItemLabel.FOOTNOTE, note, parent=table)
    assert [b.text for b in _blocks(document, None)] == [
        "Certifieringar\n"
        "Onlinetjänst | SSAE 18 SOC 1 Typ II | SSAE 18 SOC 2 Typ II\n"
        "Office 365-tjänster | Ja | Ja\n"
        "Microsoft Azure Core Services |  | Varierar*",
        note,
    ]


def test_word_headers_and_footers_are_page_furniture() -> None:
    document = DoclingDocument(name="pub")
    document.add_text(DocItemLabel.TEXT, "Sida 3 (13)", content_layer=ContentLayer.FURNITURE)
    document.add_text(DocItemLabel.TEXT, "Parterna ska ange varsin kontaktperson.")
    assert [(b.kind, b.text) for b in _blocks(document, None)] == [
        (BlockKind.PAGE_HEADER, "Sida 3 (13)"),
        (BlockKind.TEXT, "Parterna ska ange varsin kontaktperson."),
    ]
